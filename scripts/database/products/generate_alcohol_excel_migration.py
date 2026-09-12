#!/usr/bin/env python3
"""Generate SQL to migrate alcohol products from inventory Excel into Manus tax model.

Usage:
  python scripts/database/products/generate_alcohol_excel_migration.py \\
    --xlsx "l:/Bibioteca/Downloads/Reporte de Inventario (2).xlsx" \\
    --out scripts/database/products/2026_09_12_migrate_alcohol_from_excel.sql
"""

from __future__ import annotations

import argparse
import re
import uuid
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl required: pip install openpyxl") from exc

TENANT = "00000000-0000-0000-0000-000000000001"
CAT_LICORES = "d62a28a9-32ac-4df5-ad91-142c14b0edd8"
SUB_WHISKY = "7fd6dac4-6a98-4cda-822f-332ff15c791e"
SUB_AGUARDIENTE = "59469cae-477c-4263-b086-3883c987a82b"
SUB_CERVEZA = "51f91e47-5657-4a27-8b62-87b644c2912e"
TAX_CAT_LIQUOR = "16000000-0000-0000-0000-000000000010"
TAX_CAT_BEER = "16000000-0000-0000-0000-000000000014"

PACK_RE = re.compile(
    r"\b(CANASTA|CAJA|PACK|POR\s*24|POR\s*30|X24|X30|24\s*PACK)\b",
    re.I,
)
ML_RE = re.compile(
    r"(?:(\d+(?:[.,]\d+)?)\s*(?:LT|L)\b)|(?:(\d+)\s*ML\b)|(?:\b(\d{3,4})\b)",
    re.I,
)

LIQUOR_NAME_RE = re.compile(
    r"whisky|whiskey|buchan|johnnie|old\s*parr|jack\s*daniel|chivas|glenlivet|"
    r"aguardiente|ron\b|baileys|smirnoff|tequila|vodka|licor",
    re.I,
)
BEER_NAME_RE = re.compile(
    r"cerveza|aguila|poker|corona|coronita|club\s*colombia|coste|heineken|"
    r"stella|michelob|budweiser|\bred\b|bacana",
    re.I,
)


def sql_escape(value: str) -> str:
    return value.replace("'", "''")


def parse_ml(name: str) -> float | None:
    text = name.upper().replace(",", ".")
    match = ML_RE.search(text)
    if not match:
        return None
    if match.group(1):
        return float(match.group(1)) * 1000.0
    if match.group(2):
        return float(match.group(2))
    raw = float(match.group(3))
    if raw in (210, 269, 275, 310, 320, 330, 375, 500, 700, 750, 1000, 1050, 1500, 1750):
        return raw
    return None


def is_pack(name: str) -> bool:
    return bool(PACK_RE.search(name))


def classify(name: str, category: str) -> str | None:
    cat = (category or "").strip().upper()
    if is_pack(name):
        return "pack"
    if cat == "CERVEZA" or BEER_NAME_RE.search(name):
        if LIQUOR_NAME_RE.search(name) and "CERVEZA" not in name.upper() and "CORONA" not in name.upper():
            return "liquor"
        return "beer"
    if cat == "LICORES" or LIQUOR_NAME_RE.search(name):
        return "liquor"
    return None


def default_degree(kind: str, name: str) -> float:
    upper = name.upper()
    if kind == "beer":
        return 4.5
    if "WHISK" in upper or "BUCHAN" in upper or "JACK" in upper or "CHIVAS" in upper:
        return 40.0
    if "RON" in upper:
        return 35.0
    if "BAILEYS" in upper or "CREMA" in upper:
        return 17.0
    return 29.0


def default_subcategory(kind: str, name: str) -> str:
    upper = name.upper()
    if kind == "beer":
        return SUB_CERVEZA
    if "WHISK" in upper or "BUCHAN" in upper or "JACK" in upper or "CHIVAS" in upper:
        return SUB_WHISKY
    return SUB_AGUARDIENTE


def sku_from_code(code: str, name: str) -> str:
    base = re.sub(r"[^A-Za-z0-9]+", "-", code.strip())[:40].strip("-").upper()
    if not base:
        base = re.sub(r"[^A-Za-z0-9]+", "-", name)[:24].strip("-").upper()
    return f"XLS-{base}"[:64]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=3, values_only=True))

    liquor_units: list[dict] = []
    beer_units: list[dict] = []
    skipped_packs = 0

    for row in rows:
        if not row or not row[0]:
            continue
        code = str(row[0]).strip()
        name = str(row[1] or "").strip()
        category = str(row[6] or "").strip()
        try:
            sale_price = float(str(row[8]).replace(",", "")) if row[8] not in (None, "") else None
        except ValueError:
            sale_price = None
        try:
            cost_price = float(str(row[7]).replace(",", "")) if row[7] not in (None, "") else None
        except ValueError:
            cost_price = None

        kind = classify(name, category)
        if kind is None:
            continue
        if kind == "pack":
            skipped_packs += 1
            continue

        ml = parse_ml(name)
        if kind == "beer" and ml is None:
            ml = 330.0
        if kind == "liquor" and ml is None:
            ml = 750.0

        item = {
            "code": code,
            "name": name,
            "sale_price": sale_price or 0,
            "cost_price": cost_price or 0,
            "ml": ml,
            "degree": default_degree(kind, name),
            "subcategory": default_subcategory(kind, name),
            "sku": sku_from_code(code, name),
        }
        if kind == "liquor":
            liquor_units.append(item)
        else:
            beer_units.append(item)

    lines: list[str] = []
    lines.append("-- Auto-generated from inventory Excel. Do not hand-edit.")
    lines.append("-- DANE provisional = sale price for liquors.")
    lines.append("DO $excel$")
    lines.append("DECLARE")
    lines.append(f"  v_tenant UUID := '{TENANT}';")
    lines.append("  v_iva19 UUID;")
    lines.append("  v_iva5 UUID;")
    lines.append("  v_icl UUID;")
    lines.append("  v_adv UUID;")
    lines.append("  v_beer UUID;")
    lines.append("  v_product UUID;")
    lines.append("  v_matched UUID;")
    lines.append("BEGIN")
    lines.append(
        "  SELECT id INTO v_iva19 FROM taxes WHERE tenant_id = v_tenant AND UPPER(BTRIM(name)) = 'IVA 19%' LIMIT 1;"
    )
    lines.append(
        "  SELECT id INTO v_iva5 FROM taxes WHERE tenant_id = v_tenant AND UPPER(BTRIM(name)) = 'IVA 5%' LIMIT 1;"
    )
    lines.append(
        "  SELECT id INTO v_icl FROM taxes WHERE tenant_id = v_tenant AND UPPER(BTRIM(name)) = UPPER('Impuesto al consumo de licores') LIMIT 1;"
    )
    lines.append(
        "  SELECT t.id INTO v_adv FROM taxes t JOIN tax_types tt ON tt.id = t.tax_type_id"
        " WHERE t.tenant_id = v_tenant AND tt.code = 'AD_VALOREM' LIMIT 1;"
    )
    lines.append(
        "  SELECT id INTO v_beer FROM taxes WHERE tenant_id = v_tenant AND UPPER(BTRIM(name)) = UPPER('Impuesto al consumo de cervezas y refajos') LIMIT 1;"
    )
    lines.append(
        "  IF v_iva19 IS NULL OR v_iva5 IS NULL OR v_icl IS NULL OR v_adv IS NULL OR v_beer IS NULL THEN"
    )
    lines.append("    RAISE EXCEPTION 'Missing required tenant taxes';")
    lines.append("  END IF;")

    def emit_liquor(item: dict) -> None:
        pid = str(uuid.uuid5(uuid.NAMESPACE_URL, f"manus-alcohol:{item['code']}"))
        name = sql_escape(item["name"])
        sku = sql_escape(item["sku"])
        code = sql_escape(item["code"])
        lines.append("")
        lines.append(f"  -- LIQUOR {code} / {name}")
        lines.append("  v_product := NULL;")
        lines.append("  SELECT p.id INTO v_matched")
        lines.append("  FROM products p")
        lines.append("  WHERE p.tenant_id = v_tenant")
        lines.append(
            f"    AND (UPPER(BTRIM(p.sku)) = UPPER(BTRIM('{sku}'))"
        )
        lines.append(
            f"      OR UPPER(BTRIM(p.name)) = UPPER(BTRIM('{name}'))"
        )
        lines.append(
            f"      OR UPPER(BTRIM(p.name)) LIKE '%' || UPPER(BTRIM('{name}')) || '%')"
        )
        lines.append("  LIMIT 1;")
        lines.append("  IF v_matched IS NOT NULL THEN")
        lines.append("    v_product := v_matched;")
        lines.append("  ELSIF NOT EXISTS (")
        lines.append("    SELECT 1 FROM products")
        lines.append(
            f"    WHERE tenant_id = v_tenant AND UPPER(BTRIM(sku)) = UPPER(BTRIM('{sku}'))"
        )
        lines.append("  ) THEN")
        lines.append(f"    v_product := '{pid}'::uuid;")
        lines.append("    INSERT INTO products (")
        lines.append(
            "      id, tenant_id, unit_id, tax_id, name, description, sku, price, cost,"
        )
        lines.append(
            "      price_with_tax, price_without_tax, is_active, requires_lot,"
        )
        lines.append(
            "      operational_status, sale_type, measurement_unit,"
        )
        lines.append(
            "      category_id, subcategory_id, created_at, updated_at"
        )
        lines.append("    ) VALUES (")
        lines.append(
            f"      v_product, v_tenant, '10000000-0000-0000-0000-000000000001'::uuid, v_iva5, '{name}', '{name}', '{sku}',"
        )
        lines.append(
            f"      {item['sale_price']:.2f}, {item['cost_price']:.2f},"
        )
        lines.append(
            f"      {item['sale_price']:.2f}, {item['sale_price']:.2f}, TRUE, FALSE,"
        )
        lines.append("      'ACTIVE', 'UNIT', 'UND',")
        lines.append(
            f"      '{CAT_LICORES}'::uuid, '{item['subcategory']}'::uuid, NOW(), NOW()"
        )
        lines.append("    );")
        lines.append("  END IF;")
        lines.append("  IF v_product IS NOT NULL THEN")
        lines.append("    PERFORM public.prc_replace_product_taxes(")
        lines.append("      v_tenant, v_product,")
        lines.append("      jsonb_build_array(")
        lines.append(
            "        jsonb_build_object('tax_id', v_icl, 'calculation_order', 10),"
        )
        lines.append(
            "        jsonb_build_object('tax_id', v_adv, 'calculation_order', 20),"
        )
        lines.append(
            "        jsonb_build_object('tax_id', v_iva5, 'calculation_order', 100)"
        )
        lines.append("      )")
        lines.append("    );")
        lines.append("    PERFORM public.prc_upsert_product_tax_profile(")
        lines.append("      v_tenant, v_product,")
        lines.append("      jsonb_build_object(")
        lines.append(
            f"        'tax_product_category_id', '{TAX_CAT_LIQUOR}'::uuid,"
        )
        lines.append(f"        'alcohol_degree', {item['degree']:.3f},")
        lines.append(f"        'net_volume_ml', {item['ml']:.3f},")
        lines.append(
            f"        'dane_certified_retail_price', {item['sale_price']:.2f},"
        )
        lines.append("        'dane_price_effective_from', CURRENT_DATE::text,")
        lines.append("        'dane_price_effective_to', NULL")
        lines.append("      )")
        lines.append("    );")
        lines.append(
            "    UPDATE products SET tax_id = v_iva5, updated_at = NOW()"
        )
        lines.append("    WHERE id = v_product AND tenant_id = v_tenant;")
        lines.append("  END IF;")

    def emit_beer(item: dict) -> None:
        name = sql_escape(item["name"])
        sku = sql_escape(item["sku"])
        code = sql_escape(item["code"])
        lines.append("")
        lines.append(f"  -- BEER match-only {code} / {name}")
        lines.append("  SELECT p.id INTO v_matched")
        lines.append("  FROM products p")
        lines.append("  WHERE p.tenant_id = v_tenant")
        lines.append(
            f"    AND (UPPER(BTRIM(p.sku)) = UPPER(BTRIM('{sku}')) OR UPPER(BTRIM(p.name)) = UPPER(BTRIM('{name}')))"
        )
        lines.append("  LIMIT 1;")
        lines.append("  IF v_matched IS NOT NULL THEN")
        lines.append("    PERFORM public.prc_replace_product_taxes(")
        lines.append("      v_tenant, v_matched,")
        lines.append("      jsonb_build_array(")
        lines.append(
            "        jsonb_build_object('tax_id', v_beer, 'calculation_order', 10),"
        )
        lines.append(
            "        jsonb_build_object('tax_id', v_iva19, 'calculation_order', 100)"
        )
        lines.append("      )")
        lines.append("    );")
        lines.append("    PERFORM public.prc_upsert_product_tax_profile(")
        lines.append("      v_tenant, v_matched,")
        lines.append("      jsonb_build_object(")
        lines.append(
            f"        'tax_product_category_id', '{TAX_CAT_BEER}'::uuid,"
        )
        lines.append(f"        'alcohol_degree', {item['degree']:.3f},")
        lines.append(f"        'net_volume_ml', {item['ml']:.3f},")
        lines.append("        'dane_certified_retail_price', NULL,")
        lines.append("        'dane_price_effective_from', NULL,")
        lines.append("        'dane_price_effective_to', NULL")
        lines.append("      )")
        lines.append("    );")
        lines.append(
            "    UPDATE products SET tax_id = v_iva19, updated_at = NOW()"
        )
        lines.append("    WHERE id = v_matched AND tenant_id = v_tenant;")
        lines.append("  END IF;")

    for item in liquor_units:
        emit_liquor(item)
    for item in beer_units:
        emit_beer(item)

    lines.append("END")
    lines.append("$excel$;")
    lines.append("")
    lines.append(
        f"-- Generated liquor units: {len(liquor_units)}; beer match-only: {len(beer_units)}; skipped packs: {skipped_packs}"
    )

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(
        f"Wrote {out} liquors={len(liquor_units)} beers={len(beer_units)} packs_skipped={skipped_packs}"
    )


if __name__ == "__main__":
    main()
